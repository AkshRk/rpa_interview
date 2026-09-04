import os

import openpyxl
from openpyxl.styles import Font

import config

HEADERS = ["Quote", "Author", "Tags", "Length"]


def save_rows(rows, filename):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quotes"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r[0], r[1], r[2], str(len(r[0]))])
        wb.save(config.OUTPUT_DIR + "\\" + filename)
    return True


def append_to_master(rows, master_file):
    if os.path.exists(master_file) == True:
        wb = openpyxl.load_workbook(master_file)
        ws = wb["Sheet1"]
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(master_file)


def read_existing_quotes(master_file):
    wb = openpyxl.load_workbook(master_file)
    ws = wb.active
    seen = []
    for row in ws.iter_rows(min_row=2):
        seen.append(row[0].value)
    return seen


def remove_duplicates(rows, seen):
    for r in rows:
        if r[0] in seen:
            rows.remove(r)
    return rows
